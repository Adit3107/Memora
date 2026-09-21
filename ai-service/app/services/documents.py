from __future__ import annotations

import csv
import io
from pathlib import Path

from app.models.extraction import ExtractionResponse, PageResult
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

MAX_FILE_BYTES = 50 * 1024 * 1024


def extract_document(data: bytes, filename: str, content_type: str) -> ExtractionResponse:
    if not data:
        return _failure("document", "uploaded file is empty")
    if len(data) > MAX_FILE_BYTES:
        return _failure("document", "uploaded file is too large")

    detected = _detect_document_type(filename, content_type)
    try:
        if detected == "pdf":
            return _extract_pdf(data, filename)
        if detected == "docx":
            return _extract_docx(data, filename)
        if detected == "pptx":
            return _extract_pptx(data, filename)
        if detected == "txt":
            return _extract_txt(data, filename)
        if detected == "csv":
            return _extract_csv(data, filename)
        if detected == "xlsx":
            return _extract_xlsx(data, filename)
    except Exception:
        return _failure(detected, f"could not extract text from {detected}")

    return _failure("document", "unsupported document type")


def _extract_pdf(data: bytes, filename: str) -> ExtractionResponse:
    reader = PdfReader(io.BytesIO(data))
    pages: list[PageResult] = []
    parts: list[str] = []

    for index, page in enumerate(reader.pages, start=1):
        text = _clean_text(page.extract_text() or "")
        if text:
            pages.append(PageResult(index=index, text=text))
            parts.append(text)

    return _success("pdf", filename, "\n\n".join(parts), {"page_count": str(len(reader.pages))}, pages)


def _extract_docx(data: bytes, filename: str) -> ExtractionResponse:
    document = Document(io.BytesIO(data))
    parts = [_clean_text(paragraph.text) for paragraph in document.paragraphs]
    text = "\n\n".join(part for part in parts if part)
    return _success("docx", filename, text, {"paragraph_count": str(len(document.paragraphs))}, [])


def _extract_pptx(data: bytes, filename: str) -> ExtractionResponse:
    presentation = Presentation(io.BytesIO(data))
    pages: list[PageResult] = []
    parts: list[str] = []

    for index, slide in enumerate(presentation.slides, start=1):
        slide_parts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = _clean_text(shape.text)
                if text:
                    slide_parts.append(text)
        slide_text = "\n".join(slide_parts)
        if slide_text:
            pages.append(PageResult(index=index, text=slide_text))
            parts.append(f"Slide {index}: {slide_text}")

    return _success("pptx", filename, "\n\n".join(parts), {"slide_count": str(len(presentation.slides))}, pages)


def _extract_txt(data: bytes, filename: str) -> ExtractionResponse:
    text = data.decode("utf-8", errors="replace")
    return _success("txt", filename, text, {}, [])


def _extract_csv(data: bytes, filename: str) -> ExtractionResponse:
    text = data.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    lines: list[str] = []
    row_count = 0
    for row_index, row in enumerate(reader, start=1):
        row_count += 1
        cells = [f"Column {column_index}: {_clean_text(value)}" for column_index, value in enumerate(row, start=1) if _clean_text(value)]
        if cells:
            lines.append(f"Row {row_index}: {'; '.join(cells)}")
    return _success("csv", filename, "\n\n".join(lines), {"row_count": str(row_count)}, [])


def _extract_xlsx(data: bytes, filename: str) -> ExtractionResponse:
    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    pages: list[PageResult] = []
    parts: list[str] = []

    for index, sheet in enumerate(workbook.worksheets, start=1):
        rows: list[str] = []
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            cells = [f"Column {column_index}: {_clean_text(str(value))}" for column_index, value in enumerate(row, start=1) if value is not None and _clean_text(str(value))]
            if cells:
                rows.append(f"Row {row_index}: {'; '.join(cells)}")
        sheet_text = "\n".join(rows)
        if sheet_text:
            pages.append(PageResult(index=index, text=sheet_text))
            parts.append(f"Sheet {index} ({sheet.title}): {sheet_text}")

    return _success("xlsx", filename, "\n\n".join(parts), {"sheet_count": str(len(workbook.worksheets))}, pages)


def _detect_document_type(filename: str, content_type: str) -> str:
    media_type = content_type.split(";", 1)[0].strip().lower()
    extension = Path(filename).suffix.lower()

    if media_type == "application/pdf" or extension == ".pdf":
        return "pdf"
    if media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document" or extension == ".docx":
        return "docx"
    if media_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation" or extension == ".pptx":
        return "pptx"
    if media_type == "text/plain" or extension == ".txt":
        return "txt"
    if media_type in {"text/csv", "application/csv"} or extension == ".csv":
        return "csv"
    if media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" or extension == ".xlsx":
        return "xlsx"
    return "document"


def _success(content_type: str, filename: str, text: str, metadata: dict[str, str], pages: list[PageResult]) -> ExtractionResponse:
    text = _clean_text(text)
    if not text and not pages:
        return _failure(content_type, "extracted content is empty")
    return ExtractionResponse(
        success=True,
        content_type=content_type,
        title=filename,
        text=text,
        metadata=metadata,
        pages=pages,
    )


def _failure(content_type: str, message: str) -> ExtractionResponse:
    return ExtractionResponse(success=False, content_type=content_type, error=message)


def _clean_text(value: str) -> str:
    return " ".join(value.replace("\x00", "").split())


# Why this file exists:
# Python has mature document parsing libraries. Go calls this service only when
# Python is the better tool for the format or required detail level; Go still
# cleans, chunks, stores, and owns application state.
